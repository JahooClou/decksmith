"""Image guide: every picture in the deck with exact pixel sizes, topic and a Pexels pick.

    py -3.14 image_guide.py deck.pptx [--picks picks.json] [--api] [--out image-guide.md]

Reads picture slots and placed pictures that carry decksmith metadata
(images.add_slot / images.place), plus any other picture in the deck, and writes
a markdown guide and a JSON twin. For each image:
  - topic (the alt text) and the Pexels search query
  - every placement: slide, frame in inches, exact px at 1920x1080 and at
    3840x2160, aspect and orientation
  - the size the source must have (largest placement at 4K)
  - the Pexels proposition: photo page, photographer, original size, and an
    exact-size crop URL per placement (Pexels' CDN serves any w x h)
  - what is in the deck now: empty slot, embedded px, sharp enough or not
picks.json maps key -> {pexels_id, title, photographer, photographer_url, width, height}.
--api fills missing picks from the Pexels API (needs PEXELS_API_KEY; free key at
pexels.com/api). Without a key, find the photo on pexels.com and write picks.json.
"""
from __future__ import annotations

import argparse
import io
import json
import os
import sys
import urllib.parse
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from pptx import Presentation  # noqa: E402
from images import meta_of, orientation, pexels_crop_url, pexels_page, pexels_search, px_for  # noqa: E402

LICENCE = ("Pexels licence: free to use and modify, attribution not required but kind; do not sell "
           "unaltered copies or imply endorsement by people or brands shown. https://www.pexels.com/license/")


def api(path, key):
    req = urllib.request.Request("https://api.pexels.com/v1/" + path, headers={"Authorization": key})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read())


def pick_from_api(query, need_w, need_h, orient, key):
    res = api("search?" + urllib.parse.urlencode({"query": query, "orientation": orient, "per_page": 15}), key)
    for p in res.get("photos", []):
        if p["width"] >= need_w and p["height"] >= need_h:
            return dict(pexels_id=p["id"], title=p.get("alt") or "", photographer=p["photographer"],
                        photographer_url=p["photographer_url"], width=p["width"], height=p["height"])
    return None


def scan(prs):
    images = {}
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            meta = meta_of(sh)
            is_pic = sh.shape_type == 13
            if meta is None and not is_pic:
                continue
            if meta is None:
                meta = {"key": f"untracked-{si}-{sh.shape_id}", "alt": "", "query": ""}
            emb = None
            if is_pic:
                try:
                    from PIL import Image
                    with Image.open(io.BytesIO(sh.image.blob)) as im:
                        emb = im.size
                    # visible part after cropping
                    cw = 1 - (sh.crop_left + sh.crop_right)
                    ch = 1 - (sh.crop_top + sh.crop_bottom)
                    emb = (round(emb[0] * cw), round(emb[1] * ch))
                except Exception:  # noqa: BLE001
                    pass
            w1, h1 = px_for(prs, sh, 1920)
            w4, h4 = px_for(prs, sh, 3840)
            img = images.setdefault(meta["key"], dict(key=meta["key"], topic=meta.get("alt", ""),
                                                      query=meta.get("query", ""), pexels_id=meta.get("pexels"),
                                                      placements=[]))
            img["placements"].append(dict(
                slide=si, shape=sh.name, inches=(round(sh.width / 914400, 3), round(sh.height / 914400, 3)),
                px_1080p=(w1, h1), px_4k=(w4, h4), aspect=round(w1 / h1, 3), orientation=orientation(w1, h1),
                embedded_px=emb, status=("empty slot" if not is_pic else
                                         ("sharp at 4K" if emb and emb[0] >= w4 * 0.98 else
                                          ("sharp at 1080p only" if emb and emb[0] >= w1 * 0.98 else "TOO SMALL")))))
    return images


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--picks")
    ap.add_argument("--api", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()
    pptx = Path(a.pptx)
    prs = Presentation(str(pptx))
    images = scan(prs)
    picks = json.loads(Path(a.picks).read_text(encoding="utf-8")) if a.picks and Path(a.picks).exists() else {}
    key = os.environ.get("PEXELS_API_KEY")
    for k, img in images.items():
        need_w = max(p["px_4k"][0] for p in img["placements"])
        need_h = max(p["px_4k"][1] for p in img["placements"])
        img["source_min_px"] = (need_w, need_h)
        pick = picks.get(k)
        if not pick and img.get("pexels_id"):
            pick = {"pexels_id": img["pexels_id"]}
        if a.api and key:
            if pick and "width" not in pick:
                p = api(f"photos/{pick['pexels_id']}", key)
                pick.update(title=p.get("alt") or "", photographer=p["photographer"],
                            photographer_url=p["photographer_url"], width=p["width"], height=p["height"])
            elif not pick and img["query"]:
                pick = pick_from_api(img["query"], need_w, need_h, img["placements"][0]["orientation"], key)
        if pick:
            pid = pick["pexels_id"]
            pick["page_url"] = pick.get("page_url") or pexels_page(pid)
            pick["crop_urls"] = {f"slide {p['slide']}": pexels_crop_url(pid, *p["px_4k"]) for p in img["placements"]}
            if "width" in pick:
                pick["big_enough"] = pick["width"] >= need_w and pick["height"] >= need_h
        img["pexels"] = pick
        img["alternatives"] = pexels_search(img["query"], img["placements"][0]["orientation"]) if img["query"] else None
    out = Path(a.out) if a.out else pptx.with_name(pptx.stem + "-image-guide.md")
    lines = [f"# Image guide: {pptx.name}", "",
             "Pixel sizes are exact for the frame each picture fills. 1080p = the deck shown 1920 x 1080 "
             f"({1920 / (prs.slide_width / 914400):.0f} px per slide inch); 4K = 3840 x 2160. Source photos "
             "should meet the 4K size so they stay sharp on large screens and survive a crop.", ""]
    if not images:
        lines.append("No pictures in this deck.")
    for k, img in images.items():
        lines += [f"## {k}", "", f"- **Topic:** {img['topic'] or '(no alt text: add one)'}",
                  f"- **Pexels query:** `{img['query']}`" if img["query"] else "- **Pexels query:** (none)",
                  f"- **Source must be at least:** {img['source_min_px'][0]} x {img['source_min_px'][1]} px", "",
                  "| Slide | Frame (in) | Exact px at 1080p | Exact px at 4K | Aspect | In the deck now |",
                  "| --- | --- | --- | --- | --- | --- |"]
        for p in img["placements"]:
            emb = f"{p['embedded_px'][0]} x {p['embedded_px'][1]} px, {p['status']}" if p["embedded_px"] else p["status"]
            lines.append(f"| {p['slide']} | {p['inches'][0]:g} x {p['inches'][1]:g} | {p['px_1080p'][0]} x {p['px_1080p'][1]} "
                         f"| {p['px_4k'][0]} x {p['px_4k'][1]} | {p['aspect']:g} ({p['orientation']}) | {emb} |")
        if len(img["placements"]) > 1:
            lines.append("Used in more than one frame: download **one** file at the largest size and place it in every "
                         "frame (place() crops it). Morph only blends a picture smoothly into the same picture.")
        lines.append("")
        pk = img["pexels"]
        if pk:
            lines.append(f"**Pexels proposition:** [{pk.get('title') or 'photo ' + str(pk['pexels_id'])}]({pk['page_url']})"
                         + (f" by [{pk['photographer']}]({pk.get('photographer_url', '')})" if pk.get("photographer") else ""))
            if "width" in pk:
                lines.append(f"- Original: {pk['width']} x {pk['height']} px, "
                             + ("large enough for every placement" if pk.get("big_enough") else "**smaller than the 4K need: pick another**"))
            for where, u in pk["crop_urls"].items():
                lines.append(f"- Exact crop for {where}: {u}")
        else:
            lines.append("**Pexels proposition:** none yet. Search, pick one at least the size above, and add it to picks.json.")
        if img["alternatives"]:
            lines.append(f"- Alternatives: {img['alternatives']}")
        lines.append("")
    lines += ["---", LICENCE, ""]
    out.write_text("\n".join(lines), encoding="utf-8")
    out.with_suffix(".json").write_text(json.dumps(images, indent=2), encoding="utf-8")
    print(f"{len(images)} images -> {out}")
    for k, img in images.items():
        st = ", ".join(f"s{p['slide']} {p['px_1080p'][0]}x{p['px_1080p'][1]} ({p['status']})" for p in img["placements"])
        print(f"  {k}: {st}; pexels {img['pexels']['pexels_id'] if img['pexels'] else '-'}")


if __name__ == "__main__":
    main()
