"""Picture slots, exact pixel sizes and Pexels sourcing.

Every picture in a decksmith deck starts as a *slot*: a planned frame with a
key, a topic and a Pexels search query, stored on the shape itself so the image
guide (image_guide.py) can always be regenerated from the .pptx alone.

    from images import add_slot, place, px_for, pexels_crop_url, download_pexels

    slot = add_slot(slide, prs, x=8.0, y=0, w=5.333, h=7.5, key="hero",
                    topic="Night traffic as light trails", query="light trails long exposure",
                    pexels_id=18473029)
    path = download_pexels(18473029, *px_for(prs, slot, screen=3840), "assets/hero.jpg")  # ask first
    pic = place(slide, slot, path)          # crops to fill, keeps name, alt text and metadata

Metadata lives in the shape's cNvPr attributes:
  descr  = alt text (the topic, written as a description of the picture)
  title  = "img:<key>|query:<pexels query>|pexels:<photo id>"
The shape *name* stays free for Morph ('!!hero').
"""
from __future__ import annotations

import urllib.parse
import urllib.request
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from motion import q

SCREENS = {"1080p": 1920, "4k": 3840}


def px_for(prs, shape_or_size, screen=1920):
    """Exact pixel size a frame occupies when the deck is shown `screen` px wide."""
    if hasattr(shape_or_size, "width"):
        w, h = shape_or_size.width, shape_or_size.height
    else:
        w, h = (int(Inches(v)) for v in shape_or_size)
    k = screen / prs.slide_width
    return round(w * k), round(h * k)


def meta_of(shape):
    nv = shape._element.find(".//" + q("p:cNvPr"))
    raw = nv.get("title") or ""
    if not raw.startswith("img:"):
        return None
    out = {"alt": nv.get("descr") or ""}
    for part in raw.split("|"):
        k, _, v = part.partition(":")
        out[k] = v
    out["key"] = out.pop("img")
    return out


def _set_meta(shape, key, topic, query, pexels_id=None):
    nv = shape._element.find(".//" + q("p:cNvPr"))
    nv.set("descr", topic)
    meta = f"img:{key}|query:{query}"
    if pexels_id:
        meta += f"|pexels:{pexels_id}"
    nv.set("title", meta)


def add_slot(slide, prs, x, y, w, h, *, key, topic, query, pexels_id=None, name=None, fill="E3E3DF"):
    """A planned picture frame (inches). Rendered as a quiet tinted panel until
    place() swaps the real picture in."""
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    sh.name = name or f"Picture slot {key}"
    _set_meta(sh, key, topic, query, pexels_id)
    return sh


def place(slide, slot, image_path, *, focus=(0.5, 0.5)):
    """Put image_path into the slot: same frame, crop to fill (no distortion),
    same z-order, name and metadata. focus = point of the image to keep
    centred when cropping (0..1, 0..1)."""
    from PIL import Image
    meta_nv = slot._element.find(".//" + q("p:cNvPr"))
    descr, title, name = meta_nv.get("descr"), meta_nv.get("title"), slot.name
    x, y, w, h = slot.left, slot.top, slot.width, slot.height
    pic = slide.shapes.add_picture(str(image_path), x, y, w, h)
    with Image.open(image_path) as im:
        iw, ih = im.size
    frame, img = w / h, iw / ih
    if img > frame:   # too wide: crop left/right
        keep = frame / img
        extra = 1 - keep
        l = min(max(focus[0] - keep / 2, 0), extra)
        pic.crop_left, pic.crop_right = l, extra - l
    elif img < frame:  # too tall: crop top/bottom
        keep = img / frame
        extra = 1 - keep
        t = min(max(focus[1] - keep / 2, 0), extra)
        pic.crop_top, pic.crop_bottom = t, extra - t
    nv = pic._element.find(".//" + q("p:cNvPr"))
    if descr:
        nv.set("descr", descr)
    if title:
        nv.set("title", title)
    pic.name = name
    tree = slot._element.getparent()
    tree.remove(pic._element)
    slot._element.addprevious(pic._element)
    tree.remove(slot._element)
    return pic


def pexels_page(pexels_id):
    return f"https://www.pexels.com/photo/{pexels_id}/"


def pexels_search(query, orientation=None):
    url = "https://www.pexels.com/search/" + urllib.parse.quote(query) + "/"
    return url + (f"?orientation={orientation}" if orientation else "")


def pexels_crop_url(pexels_id, w, h):
    """Pexels serves any exact size through its image CDN: fit=crop gives w x h px."""
    return (f"https://images.pexels.com/photos/{pexels_id}/pexels-photo-{pexels_id}.jpeg"
            f"?auto=compress&cs=tinysrgb&fit=crop&w={w}&h={h}")


def orientation(w, h):
    r = w / h
    return "landscape" if r > 1.15 else ("portrait" if r < 0.87 else "square")


def download_pexels(pexels_id, w, h, dest):
    """Fetch the exact crop. Downloading is an outward action: ask the user first."""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(pexels_crop_url(pexels_id, w, h), headers={"User-Agent": "decksmith/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        dest.write_bytes(r.read())
    return dest
