"""Structural and taste checks for a .pptx, run before rendering.

    py -3.14 validate.py deck.pptx [--json report.json]

ERRORS (exit 1) are things PowerPoint repairs, drops or refuses:
  part does not parse, python-pptx cannot reopen, duplicate shape ids on a
  slide, animation targets a shape id that is not on the slide, duplicate
  timing node ids, build list entries without a shape, slide children out of
  schema order, zoom pointing at a missing slide, broken relationship targets.
WARNINGS are the decksmith design rules (references/design-rules.md):
  em/en dashes, text under 12 pt, word count, too many paragraphs, missing
  speaker notes, too many typefaces, hard-coded colours off the theme, pure
  black or white, Morph slides with no '!!' or same-name partner, pictures
  shown larger than their pixels allow.
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from motion import NS, q  # noqa: E402

SLD_ORDER = ["cSld", "clrMapOvr", "transition", "timing", "extLst"]
DASHES = re.compile("[–—]")


def local(n):
    if n.tag == q("mc:AlternateContent"):
        return "transition" if n.find(".//" + q("p:transition")) is not None else "AlternateContent"
    return etree.QName(n).localname


def theme_colors(z):
    names = [n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml", n)]
    out = set()
    for n in names[:1]:
        t = etree.fromstring(z.read(n))
        for c in t.find(".//" + q("a:clrScheme")).iter(q("a:srgbClr")):
            out.add(c.get("val").upper())
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--json")
    ap.add_argument("--max-words", type=int, default=45)
    a = ap.parse_args()
    path = Path(a.pptx)
    errors, warns, info = [], [], []
    z = zipfile.ZipFile(path)
    parsed = {}
    for n in z.namelist():
        if n.endswith(".xml") or n.endswith(".rels"):
            try:
                parsed[n] = etree.fromstring(z.read(n))
            except etree.XMLSyntaxError as e:
                errors.append(f"{n}: does not parse ({e})")
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        errors.append(f"python-pptx cannot reopen the file: {e}")
        prs = None
    names = set(z.namelist())
    for n, root in parsed.items():
        if not n.endswith(".rels"):
            continue
        base = n.replace("_rels/", "").replace(".rels", "")
        folder = base.rsplit("/", 1)[0] if "/" in base else ""
        for r in root:
            if r.get("TargetMode") == "External":
                continue
            t = r.get("Target")
            full = t.lstrip("/") if t.startswith("/") else _norm(folder + "/" + t)
            if full not in names:
                errors.append(f"{n}: relationship {r.get('Id')} points to missing {full}")
    theme = theme_colors(z)
    pres = parsed.get("ppt/presentation.xml")
    rels = parsed.get("ppt/_rels/presentation.xml.rels")
    target = {r.get("Id"): r.get("Target") for r in rels}
    slide_paths = []
    sld_ids = []
    for s in pres.find(q("p:sldIdLst")):
        slide_paths.append("ppt/" + target[s.get(q("r:id"))])
        sld_ids.append(s.get("id"))
    fonts = Counter()
    off_theme = Counter()
    prev_names = set()
    for i, sp_path in enumerate(slide_paths, 1):
        sld = parsed[sp_path]
        tag = f"slide {i}"
        order = [local(c) for c in sld if local(c) in SLD_ORDER]
        if order != sorted(order, key=SLD_ORDER.index):
            errors.append(f"{tag}: children out of schema order {order}")
        ids = [x.get("id") for x in sld.find(q("p:cSld")).iter(q("p:cNvPr"))]
        # a zoom's Choice and Fallback legitimately share one id
        fallback_ids = {x.get("id") for fb in sld.iter(q("mc:Fallback")) for x in fb.iter(q("p:cNvPr"))}
        dup = [k for k, v in Counter(ids).items() if v > 1 and k not in fallback_ids]
        if dup:
            errors.append(f"{tag}: duplicate shape ids {dup}")
        idset = set(ids)
        timing = sld.find(q("p:timing"))
        if timing is not None:
            ctn_ids = [c.get("id") for c in timing.iter(q("p:cTn"))]
            d = [k for k, v in Counter(ctn_ids).items() if v > 1]
            if d:
                errors.append(f"{tag}: duplicate timing node ids {d[:5]}")
            for t in timing.iter(q("p:spTgt")):
                if t.get("spid") not in idset:
                    errors.append(f"{tag}: animation targets shape id {t.get('spid')} which is not on the slide")
            for b in timing.iter(q("p:bldP"), q("p:bldGraphic")):
                if b.get("spid") not in idset:
                    errors.append(f"{tag}: build list names shape id {b.get('spid')} which is not on the slide")
            paras = {}
            for t in timing.iter(q("p:spTgt")):
                rg = t.find(".//" + q("p:pRg"))
                if rg is not None:
                    paras.setdefault(t.get("spid"), set()).add(int(rg.get("end")))
            for spid, ends in paras.items():
                shp = _shape(sld, spid)
                n_para = len(shp.findall(".//" + q("a:p"))) if shp is not None else 0
                if max(ends) >= n_para:
                    errors.append(f"{tag}: animation targets paragraph {max(ends)} of shape {spid}, which has {n_para}")
        # Morph pairing
        names_here = {x.get("name") for x in sld.find(q("p:cSld")).iter(q("p:cNvPr"))}
        is_morph = sld.find(".//" + q("p159:morph")) is not None
        if is_morph:
            bang = {n for n in names_here if n and n.startswith("!!")}
            pairs = bang & prev_names
            if not bang:
                warns.append(f"{tag}: Morph transition but no '!!' names; Morph will pair by name/type/text only")
            elif not pairs:
                warns.append(f"{tag}: Morph '!!' names {sorted(bang)} have no partner on slide {i - 1}")
            else:
                info.append(f"{tag}: Morph pairs {sorted(pairs)}")
        prev_names = names_here
        # zoom
        for zo in sld.iter(q("pslz:sldZmObj")):
            if zo.get("sldId") not in sld_ids:
                errors.append(f"{tag}: Slide Zoom points at slide id {zo.get('sldId')} which does not exist")
        # text rules
        words = 0
        for shp in sld.find(q("p:cSld")).iter(q("p:sp")):
            body = shp.find(q("p:txBody"))
            if body is None:
                continue
            ps = body.findall(q("a:p"))
            texts = ["".join(t.text or "" for t in p.iter(q("a:t"))) for p in ps]
            nonempty = [t for t in texts if t.strip()]
            if len(nonempty) > 6:
                warns.append(f"{tag}: one text box holds {len(nonempty)} paragraphs (split it, or cut)")
            for t in nonempty:
                words += len(t.split())
                if DASHES.search(t):
                    warns.append(f"{tag}: dash character in {t[:50]!r} (rule: no em/en dashes)")
            for r in body.iter(q("a:rPr"), q("a:endParaRPr")):
                sz = r.get("sz")
                if sz and int(sz) < 1200 and r.tag == q("a:rPr"):
                    txt = (r.getparent().findtext(q("a:t")) or "").strip()
                    if txt:
                        warns.append(f"{tag}: {int(sz) / 100:g} pt text {txt[:30]!r} (under 12 pt cannot be read from the back)")
                lat = r.find(q("a:latin"))
                if lat is not None:
                    fonts[lat.get("typeface")] += 1
        if words > a.max_words:
            warns.append(f"{tag}: {words} words on the slide (limit {a.max_words}; move the rest to notes)")
        for c in sld.iter(q("a:srgbClr")):
            v = c.get("val").upper()
            if v in ("000000", "FFFFFF"):
                warns.append(f"{tag}: pure #{v} (use the theme's dk1/lt1)")
            elif v not in theme:
                off_theme[v] += 1
        # notes
        if prs is not None:
            s = prs.slides[i - 1]
            if not s.has_notes_slide or not s.notes_slide.notes_text_frame.text.strip():
                warns.append(f"{tag}: no speaker notes")
            for shp in s.shapes:
                if shp.shape_type == 13 and hasattr(shp, "image"):
                    try:
                        from PIL import Image
                        im = Image.open(io.BytesIO(shp.image.blob))
                        need_w = shp.width / 914400 * 144
                        if im.width < need_w * 0.98:
                            warns.append(f"{tag}: picture {shp.name!r} is {im.width}px wide but shown at {need_w:.0f}px on a 1920px screen")
                    except Exception:  # noqa: BLE001
                        pass
    used = {f for f in fonts if f and not f.startswith("+")}
    if len(used) > 3:
        warns.append(f"deck: {len(used)} typefaces set directly on runs {sorted(used)} (keep to display + text + mono)")
    if off_theme:
        top = ", ".join(f"#{k}x{v}" for k, v in off_theme.most_common(6))
        warns.append(f"deck: colours not in the theme: {top}")
    report = dict(file=str(path), errors=errors, warnings=warns, info=info, fonts=sorted(used))
    if a.json:
        Path(a.json).write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(f"{path.name}: {len(slide_paths)} slides, {len(errors)} errors, {len(warns)} warnings")
    for e in errors:
        print("  ERROR", e)
    for w in warns:
        print("  warn ", w)
    for x in info:
        print("  info ", x)
    sys.exit(1 if errors else 0)


def _shape(sld, spid):
    for nv in sld.iter(q("p:cNvPr")):
        if nv.get("id") == spid:
            return nv.getparent().getparent()
    return None


def _norm(p):
    parts = []
    for seg in p.split("/"):
        if seg == "..":
            if parts:
                parts.pop()
        elif seg and seg != ".":
            parts.append(seg)
    return "/".join(parts)


if __name__ == "__main__":
    main()
