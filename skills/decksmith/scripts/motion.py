"""decksmith motion layer.

python-pptx has no API for animation, transitions, Morph, zoom or sections. This
module writes that XML directly into slides built with python-pptx. Every effect
template below was copied from XML that PowerPoint 365 itself wrote for the same
effect (see references/xml-reference.md), so PowerPoint reads it back as the
named effect in its Animation Pane, not as an anonymous custom effect.

    from motion import Motion, set_transition, morph_name

    m = Motion(slide)
    m.add(title, "float", start="after", delay=0.2, dur=0.8)
    m.add(body, "fade", start="click", paragraphs="each")
    m.add(bar, "path", start="with", path=(0, 0.12))       # move down 12% of slide height
    m.on_click(button).add(panel, "wipe", direction="left")  # trigger
    m.save()
    set_transition(slide, "morph", dur=1.2)

Times are seconds. Directions name the side the object comes from (entrance)
or goes to (exit): "top", "right", "bottom", "left".
"""
from __future__ import annotations

import uuid
from dataclasses import dataclass, field
from lxml import etree

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
    "pslz": "http://schemas.microsoft.com/office/powerpoint/2016/slidezoom",
    "p166": "http://schemas.microsoft.com/office/powerpoint/2016/6/main",
}
RT_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def q(tag: str) -> str:
    pre, local = tag.split(":")
    return "{%s}%s" % (NS[pre], local)


try:  # build nodes with python-pptx's parser so they get its element classes
    from pptx.oxml import oxml_parser as _PARSER
except ImportError:  # pragma: no cover
    _PARSER = etree.XMLParser()


def el(tag, attrs=None, *children, nsmap=None):
    e = _PARSER.makeelement(q(tag), nsmap=nsmap)
    for k, v in (attrs or {}).items():
        e.set(q(k) if ":" in k else k, str(v))
    for c in children:
        if c is not None:
            e.append(c)
    return e


def ms(seconds) -> str:
    if seconds in ("indefinite", None):
        return "indefinite"
    return str(int(round(float(seconds) * 1000)))


# --------------------------------------------------------------------------- effects

SIDES = {"top": 1, "right": 2, "bottom": 4, "left": 8}
WIPE = {"top": "wipe(up)", "right": "wipe(right)", "bottom": "wipe(down)", "left": "wipe(left)"}
FLY_IN = {  # side -> (x from, y from); PowerPoint's own formulas
    "top": ("#ppt_x", "0-#ppt_h/2"), "bottom": ("#ppt_x", "1+#ppt_h/2"),
    "left": ("0-#ppt_w/2", "#ppt_y"), "right": ("1+#ppt_w/2", "#ppt_y"),
}
FLY_OUT = {
    "top": ("ppt_x", "0-ppt_h/2"), "bottom": ("ppt_x", "1+ppt_h/2"),
    "left": ("0-ppt_w/2", "ppt_y"), "right": ("1+ppt_w/2", "ppt_y"),
}

# name -> (presetClass, presetID, subtype or callable(direction))
CATALOG = {
    # entrance
    "appear": ("entr", 1, 0),
    "fade": ("entr", 10, 0),
    "fly": ("entr", 2, lambda d: SIDES[d or "bottom"]),
    "float": ("entr", 42, 0),          # Float In: fade + rise 10% of slide height
    "float-down": ("entr", 47, 0),     # Descend: fade + drop in from above
    "wipe": ("entr", 22, lambda d: SIDES[d or "bottom"]),
    "zoom": ("entr", 53, 16),          # Zoom (faded zoom from the centre)
    # emphasis
    "grow": ("emph", 6, 0),            # Grow/Shrink, amount = scale (1.5 = 150%)
    "spin": ("emph", 8, 0),            # amount = degrees (360 default)
    "transparency": ("emph", 9, 0),    # amount = opacity 0..1 held until slide ends
    "pulse": ("emph", 26, 0),
    "fill-color": ("emph", 1, 2),      # amount = "RRGGBB"
    # exit
    "disappear": ("exit", 1, 0),
    "fade-out": ("exit", 10, 0),
    "fly-out": ("exit", 2, lambda d: SIDES[d or "bottom"]),
    "float-out": ("exit", 42, 0),
    "wipe-out": ("exit", 22, lambda d: SIDES[d or "bottom"]),
    "zoom-out": ("exit", 53, 32),
    # motion path (custom), path = (dx, dy) as fractions of slide width/height,
    # or an SVG-like PowerPoint path string "M 0 0 L 0.25 0 E"
    "path": ("path", 0, 0),
}
DEFAULT_DUR = {"appear": 0.001, "disappear": 0.001, "grow": 0.6, "spin": 1.2,
               "transparency": 0.5, "pulse": 0.5, "fill-color": 0.6, "path": 1.0,
               "float": 0.8, "float-out": 0.8, "float-down": 0.8}


@dataclass
class Effect:
    spid: int
    name: str
    start: str = "click"            # click | with | after
    delay: float = 0.0
    dur: float | None = None
    direction: str | None = None
    para: tuple | None = None       # (st, end) paragraph range or None
    ease: str | None = "out"        # out | in | inout | None (linear)
    amount: object = None
    path: object = None
    repeat: float | None = None     # repeat count, or "indefinite"
    rewind: bool = False
    grp: int = 0
    is_bg: bool = False

    @property
    def klass(self):
        return "mediacall" if self.name == "_media" else CATALOG[self.name][0]

    @property
    def seconds(self):
        d = self.dur if self.dur is not None else DEFAULT_DUR.get(self.name, 0.5)
        return max(d, 0.001)

    def length(self):
        """Time from the effect's own start to its end, used for 'after previous'."""
        if self.name == "_media":
            return self.delay + (self.dur or 0)
        if self.name == "transparency":
            return self.delay + self.seconds
        r = self.repeat if isinstance(self.repeat, (int, float)) else 1
        return self.delay + self.seconds * r


class _Ids:
    def __init__(self):
        self.n = 0

    def __call__(self):
        self.n += 1
        return str(self.n)


def _tgt(e: Effect):
    sp = el("p:spTgt", {"spid": e.spid})
    if e.is_bg:
        sp.append(el("p:bg"))
    elif e.para is not None:
        sp.append(el("p:txEl", None, el("p:pRg", {"st": e.para[0], "end": e.para[1]})))
    return el("p:tgtEl", None, sp)


def _cbhvr(ids, e, dur, attrs=(), cTn_attrs=None, delay=None, additive=None, rctx=None):
    ctn = el("p:cTn", {"id": ids(), "dur": dur, **(cTn_attrs or {})})
    if delay is not None:
        ctn.append(el("p:stCondLst", None, el("p:cond", {"delay": delay})))
    b = el("p:cBhvr", {k: v for k, v in (("additive", additive), ("rctx", rctx)) if v})
    b.append(ctn)
    b.append(_tgt(e))
    if attrs:
        b.append(el("p:attrNameLst", None, *[_attrname(a) for a in attrs]))
    return b


def _attrname(a):
    x = el("p:attrName")
    x.text = a
    return x


def _strval(v):
    return el("p:strVal", {"val": v})


def _set(ids, e, attr, value, dur="1", delay="0"):
    s = el("p:set", None, _cbhvr(ids, e, dur, [attr], {"fill": "hold"}, delay=delay))
    s.append(el("p:to", None, _strval(value)))
    return s


def _anim(ids, e, attr, frm, to, dur, additive=None, hold=True):
    def val(v):
        return el("p:fltVal", {"val": v}) if isinstance(v, (int, float)) else _strval(v)
    a = el("p:anim", {"calcmode": "lin", "valueType": "num"},
           _cbhvr(ids, e, dur, [attr], {"fill": "hold"} if hold else None, additive=additive))
    a.append(el("p:tavLst", None,
                el("p:tav", {"tm": "0"}, el("p:val", None, val(frm))),
                el("p:tav", {"tm": "100000"}, el("p:val", None, val(to)))))
    return a


def _filter(ids, e, transition, flt, dur):
    return el("p:animEffect", {"transition": transition, "filter": flt}, _cbhvr(ids, e, dur))


def _behaviours(ids, e: Effect):
    d = ms(e.seconds)
    hide_at = str(max(int(d) - 1, 0))
    n = e.name
    if n == "appear":
        return [_set(ids, e, "style.visibility", "visible")]
    if n == "fade":
        return [_set(ids, e, "style.visibility", "visible"), _filter(ids, e, "in", "fade", d)]
    if n == "fly":
        x, y = FLY_IN[e.direction or "bottom"]
        return [_set(ids, e, "style.visibility", "visible"),
                _anim(ids, e, "ppt_x", x, "#ppt_x", d, additive="base"),
                _anim(ids, e, "ppt_y", y, "#ppt_y", d, additive="base")]
    if n in ("float", "float-down"):
        y0 = "#ppt_y+.1" if n == "float" else "#ppt_y-.1"
        return [_set(ids, e, "style.visibility", "visible"), _filter(ids, e, "in", "fade", d),
                _anim(ids, e, "ppt_x", "#ppt_x", "#ppt_x", d),
                _anim(ids, e, "ppt_y", y0, "#ppt_y", d)]
    if n == "wipe":
        return [_set(ids, e, "style.visibility", "visible"),
                _filter(ids, e, "in", WIPE[e.direction or "bottom"], d)]
    if n == "zoom":
        return [_set(ids, e, "style.visibility", "visible"),
                _anim(ids, e, "ppt_w", 0, "#ppt_w", d), _anim(ids, e, "ppt_h", 0, "#ppt_h", d),
                _filter(ids, e, "in", "fade", d)]
    if n == "grow":
        k = int(round(float(e.amount or 1.5) * 100000))
        return [el("p:animScale", None, _cbhvr(ids, e, d, cTn_attrs={"fill": "hold"}),
                   el("p:by", {"x": k, "y": k}))]
    if n == "spin":
        deg = int(round(float(e.amount or 360) * 60000))
        return [el("p:animRot", {"by": deg}, _cbhvr(ids, e, d, ["r"], {"fill": "hold"}))]
    if n == "transparency":
        o = str(e.amount if e.amount is not None else 0.5)
        s = el("p:set", None, _cbhvr(ids, e, "indefinite", ["style.opacity"]))
        s.append(el("p:to", None, _strval(o)))
        return [s, el("p:animEffect", {"filter": "image", "prLst": f"opacity: {o}"},
                      _cbhvr(ids, e, "indefinite", rctx="IE"))]
    if n == "pulse":
        half = str(max(int(d) // 2, 1))
        fx = el("p:animEffect", {"transition": "out", "filter": "fade"},
                _cbhvr(ids, e, d, cTn_attrs={"tmFilter": "0, 0; .2, .5; .8, .5; 1, 0"}))
        sc = el("p:animScale", None, _cbhvr(ids, e, half, cTn_attrs={"autoRev": "1", "fill": "hold"}),
                el("p:by", {"x": 105000, "y": 105000}))
        return [fx, sc]
    if n == "fill-color":
        c = el("p:animClr", {"clrSpc": "rgb", "dir": "cw"}, _cbhvr(ids, e, d, ["fillcolor"], {"fill": "hold"}))
        c.append(el("p:to", None, el("a:srgbClr", {"val": str(e.amount or "000000").lstrip("#").upper()})))
        return [c, _set(ids, e, "fill.type", "solid", dur=d, delay=None),
                _set(ids, e, "fill.on", "true", dur=d, delay=None)]
    if n == "disappear":
        return [_set(ids, e, "style.visibility", "hidden")]
    if n == "fade-out":
        return [_filter(ids, e, "out", "fade", d), _set(ids, e, "style.visibility", "hidden", delay=hide_at)]
    if n == "fly-out":
        x, y = FLY_OUT[e.direction or "bottom"]
        return [_anim(ids, e, "ppt_x", "ppt_x", x, d, additive="base", hold=False),
                _anim(ids, e, "ppt_y", "ppt_y", y, d, additive="base", hold=False),
                _set(ids, e, "style.visibility", "hidden", delay=hide_at)]
    if n == "float-out":
        return [_filter(ids, e, "out", "fade", d),
                _anim(ids, e, "ppt_x", "ppt_x", "ppt_x", d, hold=False),
                _anim(ids, e, "ppt_y", "ppt_y", "ppt_y+.1", d, hold=False),
                _set(ids, e, "style.visibility", "hidden", delay=hide_at)]
    if n == "wipe-out":
        return [_filter(ids, e, "out", WIPE[e.direction or "bottom"], d),
                _set(ids, e, "style.visibility", "hidden", delay=hide_at)]
    if n == "zoom-out":
        return [_anim(ids, e, "ppt_w", "ppt_w", 0, d, hold=False), _anim(ids, e, "ppt_h", "ppt_h", 0, d, hold=False),
                _filter(ids, e, "out", "fade", d), _set(ids, e, "style.visibility", "hidden", delay=hide_at)]
    if n == "path":
        p = e.path
        if isinstance(p, (tuple, list)):
            p = f"M 0 0 L {p[0]:g} {p[1]:g} E"
        mo = el("p:animMotion", {"origin": "layout", "path": p, "pathEditMode": "relative", "ptsTypes": ""},
                _cbhvr(ids, e, d, ["ppt_x", "ppt_y"], {"fill": "hold"}))
        return [mo]
    raise ValueError(f"unknown effect {n!r}; known: {', '.join(CATALOG)}")


def _effect_par(ids, e: Effect, node_type):
    if e.name == "_media":
        ctn = el("p:cTn", {"id": ids(), "presetID": "1", "presetClass": "mediacall", "presetSubtype": "0",
                           "fill": "hold", "nodeType": node_type},
                 el("p:stCondLst", None, el("p:cond", {"delay": ms(e.delay)})))
        dur = ms(e.dur) if e.dur else "indefinite"
        cmd = el("p:cmd", {"type": "call", "cmd": "playFrom(0.0)"}, _cbhvr(ids, e, dur, cTn_attrs={"fill": "hold"}))
        ctn.append(el("p:childTnLst", None, cmd))
        return el("p:par", None, ctn)
    klass, preset, sub = CATALOG[e.name]
    sub = sub(e.direction) if callable(sub) else sub
    attrs = {"id": ids(), "presetID": preset, "presetClass": klass, "presetSubtype": sub}
    if e.repeat is not None:
        attrs["repeatCount"] = "indefinite" if e.repeat == "indefinite" else str(int(float(e.repeat) * 1000))
    if e.rewind:
        attrs["autoRev"] = "1"
    ease = {"out": {"decel": "50000"}, "in": {"accel": "50000"},
            "inout": {"accel": "50000", "decel": "50000"}}.get(e.ease or "", {})
    if e.name == "path" and e.ease == "out":
        ease = {"accel": "50000", "decel": "50000"}  # PowerPoint's default for paths
    attrs.update(ease)
    attrs["fill"] = "hold"
    if e.name != "fill-color":
        attrs["grpId"] = e.grp
    attrs["nodeType"] = node_type
    ctn = el("p:cTn", attrs, el("p:stCondLst", None, el("p:cond", {"delay": ms(e.delay)})),
             el("p:childTnLst", None, *_behaviours(ids, e)))
    return el("p:par", None, ctn)


# --------------------------------------------------------------------------- sequences

class _Sequence:
    def __init__(self, owner, trigger_spid=None):
        self.owner = owner
        self.trigger_spid = trigger_spid
        self.effects: list[Effect] = []

    def add(self, shape, effect="fade", *, start="click", delay=0.0, dur=None, direction=None,
            paragraphs=None, ease="out", amount=None, path=None, repeat=None, rewind=False,
            stagger=None, para_start=None):
        """Queue one effect. Returns self so calls can chain.

        paragraphs: None (whole shape), "each" (one effect per paragraph, like
        PowerPoint's 'By paragraph'), or an int / (st, end) range.
        stagger: with paragraphs="each", seconds between paragraph starts; all
        paragraphs then play from one click as a cascade instead of one click each.
        """
        if effect not in CATALOG:
            raise ValueError(f"unknown effect {effect!r}; known: {', '.join(CATALOG)}")
        if start not in ("click", "with", "after"):
            raise ValueError("start must be click, with or after")
        spid = _spid(shape)
        grp = self.owner._next_grp(spid)
        if effect == "path" and isinstance(path, (tuple, list)):
            # PowerPoint plays every motion path from the shape's position on the
            # slide, not from where the previous path left it. Chain them: start
            # this path at the end of the last one on the same shape.
            x0, y0 = self.owner._path_end.get(spid, (0.0, 0.0))
            x1, y1 = x0 + path[0], y0 + path[1]
            self.owner._path_end[spid] = (x1, y1)
            path = f"M {x0:g} {y0:g} L {x1:g} {y1:g} E"
        common = dict(spid=spid, name=effect, dur=dur, direction=direction, ease=ease, amount=amount,
                      path=path, repeat=repeat, rewind=rewind, grp=grp)
        if paragraphs is None:
            self.effects.append(Effect(start=start, delay=delay, **common))
            self.owner._builds.append((spid, grp, None))
            return self
        if paragraphs == "each":
            n = len(shape.text_frame.paragraphs)
            ranges = [(i, i) for i in range(n) if shape.text_frame.paragraphs[i].text.strip()]
        elif isinstance(paragraphs, int):
            ranges = [(paragraphs, paragraphs)]
        else:
            ranges = [tuple(paragraphs)]
        for k, rg in enumerate(ranges):
            if k == 0:
                st, dl = start, delay
            elif stagger is not None:
                st, dl = "with", delay + k * stagger   # cascade: overlaps the previous
            else:
                st, dl = para_start or "click", 0.0
            self.effects.append(Effect(start=st, delay=dl, para=rg, **common))
        self.owner._builds.append((spid, grp, "p"))
        return self

    def build(self, ids, interactive=False):
        """Group effects the way PowerPoint does: click group > time group > effect."""
        groups = []  # list of [auto(bool), [ [offset, [effects]] ... ]]
        for e in self.effects:
            if e.start == "click" or not groups:
                groups.append([e.start != "click", [[0.0, [e]]]])
                continue
            subs = groups[-1][1]
            if e.start == "with":
                subs[-1][1].append(e)
            else:  # after previous
                off, effs = subs[-1]
                end = off + max(x.length() for x in effs)
                subs.append([end, [e]])
        pars = []
        for gi, (auto, subs) in enumerate(groups):
            first = gi == 0
            if interactive:
                conds = [el("p:cond", {"delay": "0" if first else "indefinite"})]
            elif auto and first:
                conds = [el("p:cond", {"delay": "indefinite"}),
                         el("p:cond", {"evt": "onBegin", "delay": "0"}, el("p:tn", {"val": "2"}))]
            elif auto:
                # 'with'/'after' cannot open a later click group; treat as click
                conds = [el("p:cond", {"delay": "indefinite"})]
            else:
                conds = [el("p:cond", {"delay": "indefinite"})]
            outer = el("p:cTn", {"id": ids(), "fill": "hold"}, el("p:stCondLst", None, *conds))
            kids = el("p:childTnLst")
            for si, (off, effs) in enumerate(subs):
                inner = el("p:cTn", {"id": ids(), "fill": "hold"},
                           el("p:stCondLst", None, el("p:cond", {"delay": ms(off)})))
                ik = el("p:childTnLst")
                for ei, e in enumerate(effs):
                    if si == 0 and ei == 0:
                        nt = "afterEffect" if (auto and first and e.start == "after") else (
                            "withEffect" if (auto and first) else "clickEffect")
                    else:
                        nt = "withEffect" if ei > 0 else "afterEffect"
                    ik.append(_effect_par(ids, e, nt))
                inner.append(ik)
                kids.append(el("p:par", None, inner))
            outer.append(kids)
            pars.append(el("p:par", None, outer))
        return pars


def _spid(shape):
    return int(shape.shape_id) if hasattr(shape, "shape_id") else int(shape)


class Motion:
    """Collects the animation of one slide and writes its <p:timing>."""

    def __init__(self, slide):
        self.slide = slide
        self.main = _Sequence(self)
        self.triggers: list[_Sequence] = []
        self.media: dict[int, dict] = {}
        self._grp: dict[int, int] = {}
        self._builds: list[tuple] = []
        self._path_end: dict[int, tuple] = {}

    def _next_grp(self, spid):
        g = self._grp.get(spid, -1) + 1
        self._grp[spid] = g
        return g

    def add(self, shape, effect="fade", **kw):
        self.main.add(shape, effect, **kw)
        return self

    def on_click(self, trigger_shape):
        """Effects that play when the presenter clicks `trigger_shape` (a trigger)."""
        s = _Sequence(self, _spid(trigger_shape))
        self.triggers.append(s)
        return s

    def play_media(self, movie, *, start="after", delay=0.0, loop=False, length=None,
                   volume=None, hide_when_stopped=False):
        """Make a python-pptx movie (shapes.add_movie) play in the main sequence.

        start="after" on the first effect = 'Start: Automatically'. length is the
        clip length in seconds (ffprobe it); without it PowerPoint plays to the end.
        """
        spid = _spid(movie)
        self.media[spid] = dict(loop=loop, volume=volume, hide=hide_when_stopped)
        self.main.effects.append(Effect(spid=spid, name="_media", start=start, delay=delay, dur=length))
        return self

    # ------------------------------------------------------------------ write
    def save(self):
        sld = self.slide._element
        old = sld.find(q("p:timing"))
        keep_media, keep_seqs = {}, []
        if old is not None:
            for node in old.iter(q("p:video"), q("p:audio")):
                sp = node.find(".//" + q("p:spTgt"))
                keep_media[int(sp.get("spid"))] = node
            for seq in old.iter(q("p:seq")):
                ctn = seq.find(q("p:cTn"))
                if ctn is not None and ctn.get("nodeType") == "interactiveSeq":
                    trig = seq.find(".//" + q("p:stCondLst") + "/" + q("p:cond") + "/" + q("p:tgtEl") + "/" + q("p:spTgt"))
                    if trig is not None and int(trig.get("spid")) in keep_media:
                        keep_seqs.append(seq)   # python-pptx's click-to-pause on movies
            sld.remove(old)
        ids = _Ids()
        root_ctn = el("p:cTn", {"id": ids(), "dur": "indefinite", "restart": "never", "nodeType": "tmRoot"})
        kids = el("p:childTnLst")
        if self.main.effects:
            main_ctn = el("p:cTn", {"id": ids(), "dur": "indefinite", "nodeType": "mainSeq"})
            main_ctn.append(el("p:childTnLst", None, *self.main.build(ids)))
            seq = el("p:seq", {"concurrent": "1", "nextAc": "seek"}, main_ctn)
            seq.append(el("p:prevCondLst", None, el("p:cond", {"evt": "onPrev", "delay": "0"},
                                                    el("p:tgtEl", None, el("p:sldTgt")))))
            seq.append(el("p:nextCondLst", None, el("p:cond", {"evt": "onNext", "delay": "0"},
                                                    el("p:tgtEl", None, el("p:sldTgt")))))
            kids.append(seq)
        for spid, node in keep_media.items():
            cm = node.find(q("p:cMediaNode"))
            opt = self.media.get(spid, {})
            if opt.get("volume") is not None:
                cm.set("vol", str(int(opt["volume"] * 100000)))
            ctn = cm.find(q("p:cTn"))
            if opt.get("loop"):
                ctn.set("repeatCount", "indefinite")
            if opt.get("hide"):
                ctn.set("display", "0")
            ctn.set("id", ids())
            kids.append(node)
        for spid, opt in self.media.items():
            if spid not in keep_media:
                kids.append(_media_node(ids, spid, opt))
        for s in keep_seqs:
            for c in s.iter(q("p:cTn")):
                c.set("id", ids())
            kids.append(s)
        for s in self.triggers:
            ctn = el("p:cTn", {"id": ids(), "restart": "whenNotActive", "fill": "hold",
                               "evtFilter": "cancelBubble", "nodeType": "interactiveSeq"})
            trig = lambda: el("p:cond", {"evt": "onClick", "delay": "0"},
                              el("p:tgtEl", None, el("p:spTgt", {"spid": s.trigger_spid})))
            ctn.append(el("p:stCondLst", None, trig()))
            ctn.append(el("p:endSync", {"evt": "end", "delay": "0"}, el("p:rtn", {"val": "all"})))
            ctn.append(el("p:childTnLst", None, *s.build(ids, interactive=True)))
            seq = el("p:seq", {"concurrent": "1", "nextAc": "seek"}, ctn)
            seq.append(el("p:nextCondLst", None, trig()))
            kids.append(seq)
        root_ctn.append(kids)
        timing = el("p:timing", None, el("p:tnLst", None, el("p:par", None, root_ctn)))
        bld = self._bldlst()
        if bld is not None:
            timing.append(bld)
        _insert_in_order(sld, timing)
        return timing

    def _bldlst(self):
        tree = self.slide.shapes._spTree
        lst = el("p:bldLst")
        for spid, grp, build in self._builds:
            elm = _shape_elm(tree, spid)
            if elm is None:
                raise ValueError(f"spid {spid} is not on this slide")
            local = etree.QName(elm).localname
            if local == "sp":
                a = {"spid": spid, "grpId": grp}
                if build:
                    a["build"] = build
                has_fill = elm.find(q("p:spPr") + "/" + q("a:noFill")) is None
                if has_fill and not build:
                    a["animBg"] = "1"
                lst.append(el("p:bldP", a))
            elif local == "graphicFrame":
                lst.append(el("p:bldGraphic", {"spid": spid, "grpId": grp}, el("p:bldAsOne")))
        return lst if len(lst) else None


def _shape_elm(tree, spid):
    for nv in tree.iter(q("p:cNvPr")):
        if nv.get("id") == str(spid):
            e = nv.getparent().getparent()
            return e
    return None


def _media_node(ids, spid, opt):
    ctn = el("p:cTn", {"id": ids(), "fill": "remove", "display": "0" if opt.get("hide") else "1"},
             el("p:stCondLst", None, el("p:cond", {"delay": "indefinite"})))
    if opt.get("loop"):
        ctn.set("repeatCount", "indefinite")
    cm = el("p:cMediaNode", {"vol": str(int((opt.get("volume") or 1) * 100000))}, ctn,
            el("p:tgtEl", None, el("p:spTgt", {"spid": spid})))
    return el("p:video", None, cm)


# --------------------------------------------------------------------------- slide order

SLD_ORDER = ["cSld", "clrMapOvr", "transition", "timing", "extLst"]


def _local(node):
    if node.tag == q("mc:AlternateContent"):
        inner = node.find(".//" + q("p:transition"))
        return "transition" if inner is not None else "AlternateContent"
    return etree.QName(node).localname


def _insert_in_order(sld, node):
    name = _local(node)
    rank = SLD_ORDER.index(name)
    for i, child in enumerate(list(sld)):
        n = _local(child)
        if n in SLD_ORDER and SLD_ORDER.index(n) > rank:
            sld.insert(i, node)
            return
    sld.append(node)


# --------------------------------------------------------------------------- transitions

SPEED = lambda s: "fast" if s <= 0.5 else ("med" if s <= 0.75 else "slow")
DIRS = {"left": "l", "right": "r", "up": "u", "down": "d"}


def set_transition(slide, kind="fade", dur=0.7, *, direction=None, option="byObject",
                   advance_after=None, on_click=True):
    """Slide transition. kind: morph, fade, push, wipe, split, cover, pull, zoom, cut, none.

    morph option: byObject (default), byWord, byChar.
    direction for push/wipe/cover/pull: left/right/up/down (where it moves from).
    advance_after: seconds before the slide advances on its own (kiosk / auto decks).
    Morph is written as mc:AlternateContent with a Fade fallback, exactly as
    PowerPoint 365 saves it; older viewers show the fade.
    """
    sld = slide._element
    for node in list(sld):
        if _local(node) == "transition":
            sld.remove(node)
    if kind == "none":
        return None
    attrs = {"spd": SPEED(dur)}
    if not on_click:
        attrs["advClick"] = "0"
    if advance_after is not None:
        attrs["advTm"] = ms(advance_after)

    def body():
        d = DIRS.get(direction or "", None)
        if kind == "fade":
            return el("p:fade")
        if kind in ("push", "wipe", "cover", "pull"):
            return el("p:" + kind, {"dir": d} if d else None)
        if kind == "split":
            return el("p:split", {"orient": "vert", "dir": "out"})
        if kind == "zoom":
            return el("p:zoom")
        if kind == "cut":
            return el("p:cut")
        raise ValueError(f"unknown transition {kind!r}")

    if kind == "morph":
        choice_body = el("p159:morph", {"option": option}, nsmap={"p159": NS["p159"]})
        fallback = el("p:fade")
        req = "p159"
    else:
        choice_body = body()
        fallback = body()
        req = "p14"
    tr_choice = el("p:transition", attrs, nsmap={"p14": NS["p14"]})
    tr_choice.set(q("p14:dur"), ms(dur))
    tr_choice.append(choice_body)
    tr_fb = el("p:transition", attrs, fallback)
    ac = el("mc:AlternateContent", None,
            el("mc:Choice", {"Requires": req}, tr_choice, nsmap={req: NS[req]}),
            el("mc:Fallback", None, tr_fb), nsmap={"mc": NS["mc"]})
    _insert_in_order(sld, ac)
    return ac


def morph_name(shape, name):
    """Give a shape a '!!' name so Morph pairs it with the same-named shape on the
    next slide even when type, text or geometry differ."""
    shape.name = name if name.startswith("!!") else "!!" + name
    return shape


# --------------------------------------------------------------------------- zoom

def add_slide_zoom(slide, target_slide, image_path, left, top, width, height,
                   return_to_parent=True, transition_dur=1.0, line_rgb="E3E3DF"):
    """Slide Zoom (Insert > Zoom > Slide Zoom) with a click-to-jump fallback picture.

    image_path is the thumbnail PowerPoint shows until it regenerates its own;
    use a render of the target slide (render.py writes them). Viewers without
    zoom support (PowerPoint 2016 and older, LibreOffice, Keynote) show the
    picture, which jumps to the target slide on click.
    """
    from pptx.util import Emu
    part = slide.part
    img_part, img_rid = part.get_or_add_image_part(image_path)
    sld_rid = part.relate_to(target_slide.part, RT_SLIDE)
    shape_id = max([int(x.get("id")) for x in slide._element.iter(q("p:cNvPr"))] + [1]) + 1
    sld_id = None
    for s in slide.part.package.presentation_part.presentation.slides._sldIdLst:
        if s.rId and slide.part.package.presentation_part.related_part(s.rId) is target_slide.part:
            sld_id = s.id
    name = f"Slide Zoom {shape_id}"
    off = {"x": int(Emu(left)), "y": int(Emu(top))}
    ext = {"cx": int(Emu(width)), "cy": int(Emu(height))}

    def sppr(prefix):
        sp = el(prefix + ":spPr", None,
                el("a:xfrm", None, el("a:off", off), el("a:ext", ext)),
                el("a:prstGeom", {"prst": "rect"}, el("a:avLst")),
                el("a:ln", {"w": "3175"}, el("a:solidFill", None, el("a:srgbClr", {"val": line_rgb}))))
        return sp

    def blip(prefix):
        return el(prefix + ":blipFill", None, el("a:blip", {"r:embed": img_rid}),
                  el("a:stretch", None, el("a:fillRect")))

    zm_pr = el("pslz:zmPr", {"id": "{%s}" % str(uuid.uuid4()).upper(),
                             "returnToParent": "1" if return_to_parent else "0",
                             "transitionDur": ms(transition_dur)},
               blip("p166"), sppr("p166"), nsmap={"p166": NS["p166"]})
    zoom_obj = el("pslz:sldZmObj", {"sldId": sld_id}, zm_pr)
    gf = el("p:graphicFrame", None,
            el("p:nvGraphicFramePr", None,
               el("p:cNvPr", {"id": shape_id, "name": name}),
               el("p:cNvGraphicFramePr", None, el("a:graphicFrameLocks", {"noGrp": "1", "noChangeAspect": "1"})),
               el("p:nvPr")),
            el("p:xfrm", None, el("a:off", off), el("a:ext", ext)),
            el("a:graphic", None, el("a:graphicData", {"uri": NS["pslz"]},
                                     el("pslz:sldZm", None, zoom_obj))))
    pic = el("p:pic", None,
             el("p:nvPicPr", None,
                el("p:cNvPr", {"id": shape_id, "name": name},
                   el("a:hlinkClick", {"r:id": sld_rid, "action": "ppaction://hlinksldjump"})),
                el("p:cNvPicPr", None, el("a:picLocks", {"noGrp": "1", "noRot": "1", "noChangeAspect": "1"})),
                el("p:nvPr")),
             blip("p"), sppr("p"))
    ac = el("mc:AlternateContent", None,
            el("mc:Choice", {"Requires": "pslz"}, gf, nsmap={"pslz": NS["pslz"]}),
            el("mc:Fallback", None, pic), nsmap={"mc": NS["mc"]})
    slide.shapes._spTree.append(ac)
    return shape_id


# --------------------------------------------------------------------------- sections

SECTION_EXT = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


def set_sections(prs, sections):
    """sections: list of (name, [slide, ...]) in deck order. Every slide must be in
    exactly one section, or PowerPoint drops the section list."""
    p = prs.part._element
    ids_by_part = {prs.part.related_part(s.rId): s.id for s in prs.slides._sldIdLst}
    ext_lst = p.find(q("p:extLst"))
    if ext_lst is None:
        ext_lst = el("p:extLst")
        p.append(ext_lst)
    for ext in list(ext_lst):
        if ext.get("uri") == SECTION_EXT:
            ext_lst.remove(ext)
    lst = el("p14:sectionLst", None, nsmap={"p14": NS["p14"]})
    for name, slides in sections:
        sec = el("p14:section", {"name": name, "id": "{%s}" % str(uuid.uuid4()).upper()})
        ids = el("p14:sldIdLst")
        for s in slides:
            ids.append(el("p14:sldId", {"id": ids_by_part[s.part]}))
        sec.append(ids)
        lst.append(sec)
    ext = el("p:ext", {"uri": SECTION_EXT}, lst)
    ext_lst.insert(0, ext)
    return lst


# --------------------------------------------------------------------------- notes / links

def set_notes(slide, text):
    """Speaker notes. Paragraphs split on blank lines."""
    tf = slide.notes_slide.notes_text_frame
    parts = [t.strip() for t in text.strip().split("\n\n")]
    tf.text = parts[0]
    for t in parts[1:]:
        tf.add_paragraph().text = t
    return tf


def link_to_slide(shape, target_slide):
    """Click on shape jumps to target slide in the show (also works in PDF export)."""
    shape.click_action.target_slide = target_slide
    return shape


def count_main_effects(slide):
    """Effects PowerPoint will list in the main sequence (for cross-checks)."""
    t = slide._element.find(q("p:timing"))
    if t is None:
        return 0
    n = 0
    for seq in t.iter(q("p:seq")):
        c = seq.find(q("p:cTn"))
        if c is not None and c.get("nodeType") == "mainSeq":
            n += sum(1 for x in c.iter(q("p:cTn")) if x.get("presetClass"))
    return n
