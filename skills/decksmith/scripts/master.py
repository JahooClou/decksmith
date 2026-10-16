"""decksmith master layer: theme colours and fonts, slide master, layouts.

python-pptx can use a template but cannot design one. This module rewrites the
theme, the slide master and the layouts of a base package so a deck starts
from its own design system instead of Office's.

    from master import new_deck, Theme, Layout, PH

    theme = Theme(name="Signal", dk1="141416", lt1="F3F3F1", dk2="4A4A4F", lt2="E3E3DF",
                  accent1="E2522B", major="Clash Display", minor="Work Sans")
    prs = new_deck(theme, layouts=DEFAULT_LAYOUTS)
    cover = layout(prs, "Cover")

All sizes in inches unless noted; slide is 13.333 x 7.5 (16:9).
"""
from __future__ import annotations

import copy
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt

from motion import NS, q, el

BASE = Path(__file__).resolve().parent.parent / "assets" / "base-16x9.pptx"
EMU = 914400
W, H = 13.333, 7.5


@dataclass
class Theme:
    name: str = "decksmith"
    dk1: str = "141416"      # text on light ground, never pure black
    lt1: str = "F3F3F1"      # ground, never pure white
    dk2: str = "4A4A4F"      # secondary text
    lt2: str = "E3E3DF"      # hairlines, quiet fills
    accent1: str = "E2522B"  # THE accent. One per deck.
    # accents 2-6 feed default chart series and SmartArt: keep them neutral or
    # tints of accent1, or a default chart drags in Office blue/orange/grey
    accent2: str = "8A8A90"
    accent3: str = "C9C9C4"
    accent4: str = "2F2F33"
    accent5: str = "F0A48E"
    accent6: str = "B33B1B"
    hlink: str | None = None
    folHlink: str | None = None
    major: str = "Clash Display"   # headings  (+mj-lt)
    minor: str = "Work Sans"       # body text (+mn-lt)
    mono: str = "IBM Plex Mono"    # numbers, code, timings (set per run, not a theme slot)


@dataclass
class TextStyle:
    size: float = 20          # pt
    color: str = "tx1"        # scheme name (tx1, tx2, accent1 ...) or hex
    font: str = "+mn-lt"      # +mj-lt / +mn-lt / a family name
    bold: bool = False
    line: float = 1.15        # line spacing multiple
    tracking: float = 0       # pt, negative tightens (spc = 100 * pt)
    space_before: float = 0   # pt
    caps: bool = False


@dataclass
class PH:
    """A placeholder on a layout. kind: title, ctrTitle, subTitle, body, pic, chart, dt, ftr, sldNum."""
    kind: str
    x: float
    y: float
    w: float
    h: float
    idx: int | None = None
    anchor: str = "t"         # t / ctr / b
    style: TextStyle | None = None
    prompt: str | None = None
    name: str | None = None


@dataclass
class Layout:
    name: str
    type: str                 # PowerPoint layout type: title, obj, titleOnly, blank, secHead, twoObj, picTx
    placeholders: list = field(default_factory=list)


MARGIN = 0.75
DEFAULT_LAYOUTS = [
    Layout("Cover", "title", [
        PH("ctrTitle", MARGIN, 3.2, 8.6, 2.6, anchor="b", style=TextStyle(size=60, font="+mj-lt", line=0.92, tracking=-1.5),
           prompt="Deck title"),
        PH("subTitle", MARGIN, 5.95, 7.4, 0.8, idx=1, style=TextStyle(size=20, color="tx2"), prompt="One-line promise"),
    ]),
    Layout("Statement", "titleOnly", [
        PH("title", MARGIN, 0.7, 11.8, 1.6, anchor="t", prompt="One idea, said plainly"),
    ]),
    Layout("Content", "obj", [
        PH("title", MARGIN, 0.7, 11.8, 1.3, prompt="Slide title"),
        PH("body", MARGIN, 2.3, 7.2, 4.2, idx=1, prompt="Supporting text"),
    ]),
    Layout("Section", "secHead", [
        PH("title", MARGIN, 2.6, 9.5, 2.0, anchor="b", style=TextStyle(size=54, font="+mj-lt", line=0.92, tracking=-1.2),
           prompt="Section"),
        PH("body", MARGIN, 4.75, 7.0, 0.9, idx=1, style=TextStyle(size=20, color="tx2"), prompt="What this part answers"),
    ]),
    Layout("Blank", "blank", []),
]


# --------------------------------------------------------------------------- helpers

def emu(inches):
    return str(int(round(inches * EMU)))


def _fill(color):
    if len(color) == 6 and all(c in "0123456789abcdefABCDEF" for c in color):
        return el("a:solidFill", None, el("a:srgbClr", {"val": color.upper()}))
    return el("a:solidFill", None, el("a:schemeClr", {"val": color}))


def _rpr(tag, st: TextStyle, lang="en-US"):
    a = {"lang": lang, "sz": str(int(st.size * 100)), "b": "1" if st.bold else "0", "kern": "1200"}
    if st.tracking:
        a["spc"] = str(int(st.tracking * 100))
    if st.caps:
        a["cap"] = "all"
    r = el(tag, a)
    r.append(_fill(st.color))
    font = st.font
    r.append(el("a:latin", {"typeface": font}))
    r.append(el("a:ea", {"typeface": "+mn-ea" if font.startswith("+") else font}))
    r.append(el("a:cs", {"typeface": "+mn-cs" if font.startswith("+") else font}))
    return r


def _ppr(tag, st: TextStyle, lvl=0, algn="l"):
    p = el(tag, {"marL": "0", "indent": "0", "algn": algn, "defTabSz": "914400", "rtl": "0",
                 "eaLnBrk": "1", "latinLnBrk": "0", "hangingPunct": "1"})
    p.append(el("a:lnSpc", None, el("a:spcPct", {"val": str(int(st.line * 100000))})))
    p.append(el("a:spcBef", None, el("a:spcPts", {"val": str(int(st.space_before * 100))})))
    p.append(el("a:buNone"))
    p.append(_rpr("a:defRPr", st))
    return p


def _xfrm(x, y, w, h):
    return el("a:xfrm", None, el("a:off", {"x": emu(x), "y": emu(y)}), el("a:ext", {"cx": emu(w), "cy": emu(h)}))


def _sp_tree():
    return el("p:spTree", None,
              el("p:nvGrpSpPr", None, el("p:cNvPr", {"id": "1", "name": ""}), el("p:cNvGrpSpPr"), el("p:nvPr")),
              el("p:grpSpPr", None, el("a:xfrm", None, el("a:off", {"x": "0", "y": "0"}), el("a:ext", {"cx": "0", "cy": "0"}),
                                       el("a:chOff", {"x": "0", "y": "0"}), el("a:chExt", {"cx": "0", "cy": "0"}))))


PH_NAMES = {"title": "Title", "ctrTitle": "Title", "subTitle": "Subtitle", "body": "Text", "pic": "Picture",
            "chart": "Chart", "dt": "Date", "ftr": "Footer", "sldNum": "Slide Number"}


def _ph_sp(sid, ph: PH, lvl_style=None, master=False):
    ph_attrs = {}
    if ph.kind not in ("body", "obj") or master:
        ph_attrs["type"] = ph.kind
    if ph.kind == "body" and not master:
        ph_attrs["type"] = "body"
    if ph.idx is not None:
        ph_attrs["idx"] = str(ph.idx)
    if ph.kind in ("dt", "ftr", "sldNum"):
        ph_attrs["sz"] = "quarter"
    name = ph.name or f"{PH_NAMES.get(ph.kind, ph.kind)} {sid - 1}"
    sp = el("p:sp", None,
            el("p:nvSpPr", None, el("p:cNvPr", {"id": str(sid), "name": name}),
               el("p:cNvSpPr", None, el("a:spLocks", {"noGrp": "1"})),
               el("p:nvPr", None, el("p:ph", ph_attrs))),
            el("p:spPr", None, _xfrm(ph.x, ph.y, ph.w, ph.h)))
    body = el("p:txBody", None, el("a:bodyPr", {"vert": "horz", "lIns": "0", "tIns": "0", "rIns": "0", "bIns": "0",
                                                 "rtlCol": "0", "anchor": ph.anchor},
                                    el("a:normAutofit")))
    lst = el("a:lstStyle")
    if ph.style is not None:
        lst.append(_ppr("a:lvl1pPr", ph.style, algn="r" if ph.kind == "sldNum" else "l"))
    elif ph.kind == "sldNum":
        lst.append(el("a:lvl1pPr", {"algn": "r"}))
    body.append(lst)
    para = el("a:p")
    if ph.kind == "sldNum":
        f = el("a:fld", {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", "type": "slidenum"},
               el("a:rPr", {"lang": "en-US"}))
        t = el("a:t")
        t.text = "‹#›"
        f.append(t)
        para.append(f)
    elif ph.prompt or master:
        r = el("a:r", None, el("a:rPr", {"lang": "en-US"}))
        t = el("a:t")
        t.text = ph.prompt or {"title": "Click to edit title", "body": "Click to edit text"}.get(ph.kind, "")
        r.append(t)
        para.append(r)
    para.append(el("a:endParaRPr", {"lang": "en-US"}))
    body.append(para)
    sp.append(body)
    return sp


# --------------------------------------------------------------------------- theme

def apply_theme(prs, t: Theme, part=None):
    """Rewrite the colour scheme, font scheme and name of the deck's theme."""
    part = part or prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    root.set("name", t.name)
    cs = root.find(".//" + q("a:clrScheme"))
    cs.set("name", t.name)
    vals = dict(dk1=t.dk1, lt1=t.lt1, dk2=t.dk2, lt2=t.lt2, accent1=t.accent1, accent2=t.accent2,
                accent3=t.accent3, accent4=t.accent4, accent5=t.accent5, accent6=t.accent6,
                hlink=t.hlink or t.accent1, folHlink=t.folHlink or t.dk2)
    for slot, hexv in vals.items():
        node = cs.find(q("a:" + slot))
        for c in list(node):
            node.remove(c)
        node.append(el("a:srgbClr", {"val": hexv.upper()}))
    fs = root.find(".//" + q("a:fontScheme"))
    fs.set("name", t.name)
    for which, face in (("majorFont", t.major), ("minorFont", t.minor)):
        f = fs.find(q("a:" + which))
        f.find(q("a:latin")).set("typeface", face)
        for sub in f.findall(q("a:font")):   # per-script overrides would fight the latin face
            f.remove(sub)
    blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if hasattr(part, "_element"):
        part._element = parse_xml(blob)
    else:
        part._blob = blob
    return part


def theme_notes(prs, t: Theme):
    """python-pptx creates the notes master with its own Calibri theme. Call this
    after the notes are written so speaker notes use the deck's fonts too."""
    return apply_theme(prs, t, prs.notes_master.part.part_related_by(RT.THEME))


# --------------------------------------------------------------------------- master

def style_master(prs, t: Theme, *, title=None, body=None, ground="bg1", slide_number=True,
                 decorate=None):
    """Rebuild the master: ground, text styles, placeholders. decorate(spTree) may
    append fixed shapes (a rule, a logo) that every layout inherits."""
    title = title or TextStyle(size=40, font="+mj-lt", line=0.95, tracking=-0.8)
    body = body or TextStyle(size=20, font="+mn-lt", color="tx1", line=1.2, space_before=8)
    m = prs.slide_master._element
    csld = m.find(q("p:cSld"))
    new = el("p:cSld", None,
             el("p:bg", None, el("p:bgPr", None, _fill(ground), el("a:effectLst"))))
    tree = _sp_tree()
    tree.append(_ph_sp(2, PH("title", MARGIN, 0.7, 11.8, 1.3), master=True))
    tree.append(_ph_sp(3, PH("body", MARGIN, 2.3, 11.8, 4.2, idx=1), master=True))
    tree.append(_ph_sp(4, PH("dt", MARGIN, 6.85, 2.5, 0.3, idx=2, style=TextStyle(size=10, color="tx2", font=t.mono)), master=True))
    tree.append(_ph_sp(5, PH("ftr", 4.4, 6.85, 4.5, 0.3, idx=3, style=TextStyle(size=10, color="tx2", font=t.mono)), master=True))
    tree.append(_ph_sp(6, PH("sldNum", W - MARGIN - 1.5, 6.85, 1.5, 0.3, idx=4,
                             style=TextStyle(size=10, color="tx2", font=t.mono)), master=True))
    if decorate:
        decorate(tree)
    new.append(tree)
    m.replace(csld, new)
    tx = m.find(q("p:txStyles"))
    styles = el("p:txStyles")
    ts = el("p:titleStyle", None, _ppr("a:lvl1pPr", title))
    styles.append(ts)
    bs = el("p:bodyStyle")
    for lvl in range(1, 10):
        st = copy.copy(body)
        if lvl > 1:
            st.size = max(body.size - 2 * (lvl - 1), 12)
            st.color = "tx2"
        p = _ppr(f"a:lvl{lvl}pPr", st)
        p.set("marL", emu(0.3 * (lvl - 1)))
        bs.append(p)
    styles.append(bs)
    os_ = el("p:otherStyle", None, _ppr("a:lvl1pPr", TextStyle(size=18)))
    styles.append(os_)
    m.replace(tx, styles)
    # remember: the slide-number placeholder only shows on slides that carry one
    return prs.slide_master


# --------------------------------------------------------------------------- layouts

def build_layouts(prs, layouts, t: Theme):
    """Replace the base layouts with `layouts`. Existing layout parts are reused
    in order and the surplus is removed."""
    master = prs.slide_master
    id_lst = master._element.find(q("p:sldLayoutIdLst"))
    entries = list(id_lst)
    rids = [e.get(q("r:id")) for e in entries]
    parts = [master.part.related_part(r) for r in rids]
    if len(layouts) > len(parts):
        raise ValueError("base template has fewer layouts than requested")
    for spec, part in zip(layouts, parts):
        root = el("p:sldLayout", {"type": spec.type, "preserve": "1"},
                  nsmap={"a": NS["a"], "r": NS["r"], "p": NS["p"]})
        csld = el("p:cSld", {"name": spec.name})
        tree = _sp_tree()
        sid = 2
        for ph in spec.placeholders:
            tree.append(_ph_sp(sid, ph))
            sid += 1
        for kind, x, idx in (("dt", MARGIN, 10), ("ftr", 4.4, 11), ("sldNum", W - MARGIN - 1.5, 12)):
            w = 1.5 if kind == "sldNum" else (2.5 if kind == "dt" else 4.5)
            tree.append(_ph_sp(sid, PH(kind, x, 6.85, w, 0.3, idx=idx)))
            sid += 1
        csld.append(tree)
        root.append(csld)
        root.append(el("p:clrMapOvr", None, el("a:masterClrMapping")))
        part._element = root
        part.__dict__.pop("slide_layout", None)
    for rid, entry in zip(rids[len(layouts):], entries[len(layouts):]):
        id_lst.remove(entry)
        master.part.drop_rel(rid)
    return [l.name for l in prs.slide_layouts]


def layout(prs, name):
    for l in prs.slide_layouts:
        if l.name == name:
            return l
    raise KeyError(f"no layout {name!r}; have {[l.name for l in prs.slide_layouts]}")


def new_deck(theme: Theme, layouts=None, author="", title="", **master_kw):
    """A fresh 16:9 deck with this theme, master and layouts, and no slides.
    author/title go into File > Info; the base template carries none."""
    prs = Presentation(str(BASE))
    cp = prs.core_properties
    cp.author, cp.last_modified_by, cp.title = author, author, title
    apply_theme(prs, theme)
    style_master(prs, theme, **master_kw)
    build_layouts(prs, layouts or DEFAULT_LAYOUTS, theme)
    return prs


def show_slide_number(slide, layout_ph_idx=12):
    """Clone the layout's slide-number placeholder onto a slide (python-pptx
    does not copy footer placeholders when it adds a slide)."""
    lay = slide.slide_layout
    for ph in lay.placeholders:
        if ph.placeholder_format.type is not None and ph.placeholder_format.idx == layout_ph_idx:
            new = copy.deepcopy(ph._element)
            ids = [int(x.get("id")) for x in slide._element.iter(q("p:cNvPr"))]
            new.find(".//" + q("p:cNvPr")).set("id", str(max(ids) + 1))
            slide.shapes._spTree.append(new)
            return new
    return None
