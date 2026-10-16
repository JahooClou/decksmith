"""decksmith demo deck: 7 slides, custom master, Morph, sequenced builds, a motion
path, a native chart, a Slide Zoom, a trigger, sections, a jump link, notes and
two Pexels photos. Also the worked example the SKILL.md points to.

    py -3.14 build_demo.py --out Z:/Claude/decksmith-demo [--fetch] [--thumbs <render dir>]

--fetch   download the Pexels photos at their exact 4K crop (ask the user first)
--thumbs  folder with slide-02.png from a first render; used as the zoom thumbnail
Design read: explainer deck for a creative team; monochrome + one vermilion;
Clash Display headings, Work Sans text, IBM Plex Mono for code and seconds;
VARIANCE 6 / MOTION 5 / DENSITY 3. Every effect has a reason in the notes.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from images import add_slot, download_pexels, place, px_for  # noqa: E402
from master import Theme, layout, new_deck, show_slide_number, theme_notes  # noqa: E402
from motion import (Motion, add_slide_zoom, link_to_slide, morph_name, set_notes,  # noqa: E402
                    set_sections, set_transition)

T = Theme(name="decksmith Signal")
INK, GROUND, GREY, HAIR, POP = T.dk1, T.lt1, T.dk2, T.lt2, T.accent1
DISPLAY, TEXT, MONO = T.major, T.minor, T.mono

PHOTOS = {
    "hero": dict(pexels_id=18473029, topic="Night traffic drawn as light trails on a curved city road",
                 query="light trails long exposure", title="Car Lights Trails on Road Curve in City",
                 photographer="King Ho", photographer_url="https://www.pexels.com/@king/", width=4016, height=6016),
    "beam": dict(pexels_id=20728625, topic="A projector beam cutting through a dark room",
                 query="projector light beam dark room", title="Light of Projector",
                 photographer="Natasha Latinovska",
                 photographer_url="https://www.pexels.com/@natasha-latinovska-1052962659/", width=4000, height=6000),
}


def rgb(h):
    return RGBColor.from_string(h)


def text(slide, x, y, w, h, paras, *, font=TEXT, size=20, color=INK, bold=False, line=1.15,
         anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, space=6, name=None, tracking=None):
    """A text box with no insets, one run per paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        para.space_before = Pt(space if i else 0)
        r = para.add_run()
        r.text = p
        f = r.font
        f.name, f.size, f.bold = font, Pt(size), bold
        f.color.rgb = rgb(color)
        if tracking:
            r._r.get_or_add_rPr().set("spc", str(int(tracking * 100)))
    return tb


def block(slide, x, y, w, h, color=POP, name=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    return sh


def set_title(slide, s, x=None, y=None, w=None, h=None, size=None):
    t = slide.shapes.title
    t.text = s
    if x is not None:
        t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(h)
    if size:
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)
    return t


def photo(slide, prs, key, x, y, w, h, assets, fetch, name):
    p = PHOTOS[key]
    slot = add_slot(slide, prs, x, y, w, h, key=key, topic=p["topic"], query=p["query"],
                    pexels_id=p["pexels_id"], name=name)
    src = assets / f"{key}.jpg"
    if fetch and not src.exists():
        # the largest placement decides the source size; hero is biggest on slide 1
        download_pexels(p["pexels_id"], *SOURCE_PX[key], src)
    if src.exists():
        return place(slide, slot, src)
    return slot


RULE_Y = 1.72
SOURCE_PX = {"hero": (1536, 2160), "beam": (1325, 2160)}   # px_for(...) at 3840 of the largest frame


def build(out: Path, fetch=False, thumbs: Path | None = None):
    assets = out / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    prs = new_deck(T)
    COVER, STATEMENT, CONTENT, BLANK = (layout(prs, n) for n in ("Cover", "Statement", "Content", "Blank"))

    # ---------------------------------------------------------------- 1 cover
    s1 = prs.slides.add_slide(COVER)
    hero1 = photo(s1, prs, "hero", 8.0, 0, 5.333, 7.5, assets, fetch, "!!hero")
    pop1 = morph_name(block(s1, 0.75, 0.75, 0.42, 0.42), "pop")
    t1 = set_title(s1, "Motion that carries meaning", 0.75, 2.9, 6.7, 2.9, size=58)
    sub1 = s1.placeholders[1]
    sub1.text = "Animated PowerPoint, built from code and checked in PowerPoint."
    sub1.left, sub1.top, sub1.width, sub1.height = Inches(0.75), Inches(6.0), Inches(6.6), Inches(0.8)
    m = Motion(s1)
    m.add(t1, "float", start="after", delay=0.3, dur=0.9)
    m.add(sub1, "fade", start="with", delay=0.7, dur=0.6)
    m.save()
    set_transition(s1, "fade", dur=0.8)
    set_notes(s1, """Open on the promise. The title floats in on its own, so the first click is yours.

The photo is the object Morph will carry into the next slide; the small vermilion square is the second one. Watch both.""")

    # ---------------------------------------------------------------- 2 morph
    s2 = prs.slides.add_slide(STATEMENT)
    photo(s2, prs, "hero", 8.6, 1.0, 3.98, 5.5, assets, fetch, "!!hero")
    morph_name(block(s2, 0.75, RULE_Y, 1.3, 0.09), "pop")
    set_title(s2, "Morph moves the eye")
    body2 = text(s2, 0.75, 2.45, 7.1, 2.4, [
        "Give an object the same name on two slides, starting with two exclamation marks.",
        "PowerPoint animates the position, size, crop and colour in between. No timeline to build."],
        size=21, color=INK, line=1.2, space=14)
    code2 = text(s2, 0.75, 5.35, 6.0, 0.5, "!!hero    !!pop", font=MONO, size=18, color=GREY)
    m = Motion(s2)
    m.add(body2, "fade", start="after", delay=0.2, dur=0.6, paragraphs="each", stagger=0.35)
    m.add(code2, "fade", start="after", delay=0.1, dur=0.5)
    m.save()
    set_transition(s2, "morph", dur=1.4)
    show_slide_number(s2)
    set_notes(s2, """This slide arrived by Morph: the photo shrank into its frame and the square stretched into the rule under the title. Nothing was keyframed.

The rule: same name on both slides, and a name that starts with two exclamation marks forces the match even when the shapes differ. The two names are on screen in mono.""")

    # ---------------------------------------------------------------- 3 builds + path
    s3 = prs.slides.add_slide(STATEMENT)
    set_title(s3, "Three kinds of motion, one job each")
    rows = [("Entrance", "Brings in the point you are about to make."),
            ("Emphasis", "Points at what matters while you talk."),
            ("Exit", "Clears what you have finished with.")]
    top, step = 2.35, 1.25
    bar = morph_name(block(s3, 0.75, top + 0.02, 0.09, 0.78), "pop")
    words, notes3 = [], []
    for i, (w, d) in enumerate(rows):
        words.append(text(s3, 1.3, top + i * step, 3.6, 0.8, w, font=DISPLAY, size=36, color=INK,
                          anchor=MSO_ANCHOR.MIDDLE, tracking=-0.5))
        notes3.append(text(s3, 5.0, top + i * step, 6.8, 0.8, d, size=20, color=GREY, anchor=MSO_ANCHOR.MIDDLE))
    m = Motion(s3)
    timeline = []          # (label, start offset in its click, duration): drives slide 4's chart
    for i in range(3):
        if i:
            m.add(bar, "path", start="click", path=(0, step / 7.5), dur=0.6)
            timeline.append((f"Marker to {rows[i][0].lower()}", 0.0, 0.6))
            m.add(words[i], "float", start="with", delay=0.15, dur=0.8)
            timeline.append((f"{rows[i][0]}: word", 0.15, 0.8))
        else:
            m.add(words[i], "float", start="click", dur=0.8)
            timeline.append((f"{rows[i][0]}: word", 0.0, 0.8))
        base = 0.15 if i else 0.0
        m.add(notes3[i], "fade", start="with", delay=base + 0.3, dur=0.6)
        timeline.append((f"{rows[i][0]}: line", base + 0.3, 0.6))
    m.add(words[1], "pulse", start="click", dur=0.5)
    timeline.append(("Emphasis: pulse", 0.0, 0.5))
    m.add(words[2], "fade-out", start="click", dur=0.5)
    m.add(notes3[2], "fade-out", start="with", dur=0.5)
    timeline.append(("Exit: word and line", 0.0, 0.5))
    m.save()
    set_transition(s3, "morph", dur=1.0)
    show_slide_number(s3)
    set_notes(s3, """Three clicks build the three rows. On each, the marker moves down by a motion path and the word floats in with it, then the line fades in a beat later: the order tells the room what to read first.

Click four pulses Emphasis. Click five takes Exit away, which is what exits are for.""")

    # ---------------------------------------------------------------- 4 chart
    s4 = prs.slides.add_slide(STATEMENT)
    set_title(s4, "Each click is a short score")
    morph_name(block(s4, 0.75, RULE_Y, 1.3, 0.09), "pop")
    cd = CategoryChartData()
    cd.categories = [t[0] for t in timeline]
    cd.add_series("Wait", [round(t[1], 2) for t in timeline])
    cd.add_series("Plays", [round(t[2], 2) for t in timeline])
    gf = s4.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.75), Inches(2.3), Inches(8.6), Inches(4.4), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.name, ch.font.size, ch.font.color.rgb = TEXT, Pt(13), rgb(GREY)
    wait, plays = ch.plots[0].series
    wait.format.fill.background()
    plays.format.fill.solid()
    plays.format.fill.fore_color.rgb = rgb(POP)
    ch.plots[0].gap_width = 60
    ca, va = ch.category_axis, ch.value_axis
    ca.reverse_order = True
    ca.format.line.fill.background()
    ca.has_major_gridlines = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(HAIR)
    va.format.line.fill.background()
    va.maximum_scale, va.major_unit = 1.2, 0.2
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '0.0" s"', False
    va.tick_labels.font.name = MONO
    side = text(s4, 9.9, 2.45, 2.7, 3.2, [
        "Seconds from each click on the previous slide.",
        "With previous overlaps. After previous waits for the end."], size=17, color=GREY, line=1.2, space=12)
    m = Motion(s4)
    m.add(gf, "wipe", start="after", delay=0.2, direction="left", dur=0.9)
    m.add(side, "fade", start="after", dur=0.5)
    m.save()
    set_transition(s4, "morph", dur=0.8)
    show_slide_number(s4)
    set_notes(s4, """This chart is the real timing of the previous slide, read from the same code that wrote it, so it cannot drift from what plays.

Vermilion bars are effects playing; the gap before each bar is its delay. A native chart, so it stays editable in PowerPoint.""")

    # ---------------------------------------------------------------- 5 zoom + trigger
    s5 = prs.slides.add_slide(STATEMENT)
    set_title(s5, "Zoom jumps in, then comes back")
    morph_name(block(s5, 0.75, RULE_Y, 1.3, 0.09), "pop")
    text(s5, 0.75, 2.45, 6.0, 1.6, ["Click the frame to zoom into slide 2.",
                                    "Then the show comes back here."], size=21, line=1.2, space=10)
    thumb = (thumbs / "slide-02.png") if thumbs else None
    if not thumb or not thumb.exists():
        from PIL import Image
        thumb = assets / "zoom-placeholder.png"
        Image.new("RGB", (1280, 720), "#" + HAIR).save(thumb)
    zx, zy, zw = 7.0, 2.45, 5.583
    add_slide_zoom(s5, s2, str(thumb), Inches(zx), Inches(zy), Inches(zw), Inches(zw * 9 / 16), line_rgb=HAIR)
    btn = text(s5, 0.75, 4.55, 3.6, 0.45, "Why does it come back?", font=MONO, size=16, color=POP, name="Rule trigger")
    rule = text(s5, 0.75, 5.15, 5.6, 0.9, "Return to zoom is on, so slide 2 hands control back to this slide.",
                size=18, color=GREY, line=1.2)
    m = Motion(s5)
    m.on_click(btn).add(rule, "wipe", direction="left", dur=0.5)
    m.save()
    set_transition(s5, "morph", dur=0.8)
    show_slide_number(s5)
    set_notes(s5, """Slide Zoom: the frame is a live link into slide 2 that zooms in and, with return to zoom on, comes back here. Older viewers see a picture that jumps to slide 2.

The vermilion line is a trigger. Click it, not the slide, to reveal the answer: useful for questions you may or may not get.""")

    # ---------------------------------------------------------------- 6 limits
    s6 = prs.slides.add_slide(STATEMENT)
    photo(s6, prs, "beam", 0, 0, 4.6, 7.5, assets, fetch, "Projector photo")
    set_title(s6, "What only PowerPoint plays", 5.3, 0.7, 7.3, 1.3)
    g1 = text(s6, 5.3, 2.55, 3.65, 3.4, ["Any viewer", "Layouts, theme colours and fonts",
                                        "Speaker notes and sections", "Charts and links"], size=19, color=GREY,
              line=1.15, space=12)
    g1.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(INK)
    g1.text_frame.paragraphs[0].runs[0].font.name = DISPLAY
    g1.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    g2 = text(s6, 9.15, 2.55, 3.65, 3.4, ["PowerPoint only", "Morph and Slide Zoom",
                                        "Animation timing and triggers", "Media that plays by itself"],
              size=19, color=GREY, line=1.2, space=12)
    g2.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(INK)
    g2.text_frame.paragraphs[0].runs[0].font.name = DISPLAY
    g2.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    under = morph_name(block(s6, 9.15, 2.3, 1.3, 0.09), "pop")
    m = Motion(s6)
    m.add(g2, "fade", start="click", dur=0.6)
    m.add(under, "fade", start="with", dur=0.6)
    m.save()
    set_transition(s6, "fade", dur=0.7)
    show_slide_number(s6)
    set_notes(s6, """Be honest about where this works. The left column survives Keynote, Google Slides, LibreOffice and a PDF. The right column needs PowerPoint 365 or 2019 and later.

So check every deck in PowerPoint before it goes in front of people; the render script does that part.""")

    # ---------------------------------------------------------------- 7 close
    s7 = prs.slides.add_slide(BLANK)
    morph_name(block(s7, 8.0, 0, 5.333, 7.5), "pop")
    close = text(s7, 0.75, 2.5, 6.6, 2.4, "Build it, render it, then watch it in PowerPoint.",
                 font=DISPLAY, size=46, line=0.95, tracking=-1.0, anchor=MSO_ANCHOR.BOTTOM)
    path7 = text(s7, 0.75, 5.15, 6.8, 0.4, "Z:\\Claude\\.claude\\skills\\decksmith", font=MONO, size=15, color=GREY)
    btn7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(5.9), Inches(2.7), Inches(0.62))
    btn7.fill.solid()
    btn7.fill.fore_color.rgb = rgb(INK)
    btn7.line.fill.background()
    btn7.shadow.inherit = False
    btn7.name = "Back to start"
    tf = btn7.text_frame
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.text = "Back to the start"
    r = tf.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.color.rgb = TEXT, Pt(16), rgb(GROUND)
    link_to_slide(btn7, s1)
    m = Motion(s7)
    m.add(close, "float", start="after", delay=0.4, dur=0.9)
    m.add(path7, "fade", start="after", dur=0.4)
    m.add(btn7, "fade", start="with", delay=0.2, dur=0.4)
    m.save()
    set_transition(s7, "morph", dur=1.6)
    set_notes(s7, """The square from slide 6 grows into the whole right side: one more Morph, and the deck closes on the colour it opened with.

The button jumps back to slide 1 in the show. The skill lives at the path on screen.""")

    set_sections(prs, [("Opening", [s1]), ("Motion", [s2, s3, s4]), ("Navigation", [s5]), ("Limits", [s6, s7])])
    theme_notes(prs, T)
    pptx = out / "decksmith-demo.pptx"
    prs.save(pptx)
    (out / "picks.json").write_text(json.dumps(PHOTOS, indent=2), encoding="utf-8")
    print(pptx)
    return pptx


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=r"Z:\Claude\decksmith-demo")
    ap.add_argument("--fetch", action="store_true")
    ap.add_argument("--thumbs")
    a = ap.parse_args()
    build(Path(a.out), a.fetch, Path(a.thumbs) if a.thumbs else None)
